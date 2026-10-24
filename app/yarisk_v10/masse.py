"""Traitement en masse : synthèse d'un lot de liasses (une ligne par liasse) et export Excel.

Chaque JSON V29d du lot est traité indépendamment (pas de consolidation entre clients) : on retient l'exercice N
de la liasse et on calcule les agrégats avec les mêmes formules que l'écran Bilan.
"""
import io
import json
import random
from datetime import datetime
from pathlib import Path

import bilan_view as bv
import fiche_client as fc

COLUMNS = [
    ("fichier", "Fichier", None), ("raison_sociale", "Raison sociale", None), ("nif", "NIF", None),
    ("exercice", "Exercice", None), ("nature", "Nature", None),
    ("ca", "Chiffre d'affaires (DZD)", "n"), ("ebe", "EBE (DZD)", "n"), ("rn", "Résultat net (DZD)", "n"),
    ("fp", "Fonds propres (DZD)", "n"), ("tb", "Total bilan (DZD)", "n"), ("dfin", "Dettes financières (DZD)", "n"),
    ("r_autonomie", "Autonomie financière", "p"), ("r_marge_ebe", "Marge d'EBE", "p"),
    ("equilibre", "Actif = Passif", None), ("ecart_bilan", "Écart actif − passif (DZD)", "n"),
    ("nb_anomalies", "Anomalies (≥ 1 000 DZD)", "i"), ("statut", "Statut", None),
]


def nature_of(doc):
    ident = doc.get("identite") or {}
    n = (ident.get("nature") or "").upper()
    if n:
        return {"FISCAL": "Fiscal", "CNRC": "CNRC", "SITUATION": "Situation"}.get(n, n.title())
    src = str((doc.get("document") or {}).get("source") or "")
    return "CNRC" if "CNRC" in src.upper() else "Fiscal"


def summarize(doc_id, doc):
    ident = doc.get("identite") or {}
    row = {"fichier": (doc.get("document") or {}).get("fichier_source") or doc_id, "raison_sociale": ident.get("entreprise"),
           "nif": ident.get("nif_15"), "exercice": ident.get("annee"), "nature": nature_of(doc), "statut": "OK"}
    try:
        cons = bv.consolidate({doc_id: doc})
        y = bv.doc_year(doc) or (max(cons["years"]) if cons["years"] else None)
        fin = fc.financials(cons)
        d = fin.get(y) or {}
        for k in ("ca", "ebe", "rn", "fp", "tb", "dfin", "r_autonomie", "r_marge_ebe"):
            row[k] = d.get(k)
        row["exercice"] = y
        mc = bv.multi_controls(cons)
        cr = [c for c in mc["cross"] if c["annee"] == y and c["controle"].startswith("Total actif")]
        if cr:
            row["ecart_bilan"] = cr[0]["ecart"]
            row["equilibre"] = "OK" if cr[0]["niveau"] == "OK" else "ÉCART"
        row["nb_anomalies"] = sum(1 for a in mc["alerts"] if a["annee"] == y) + sum(1 for c in mc["cross"] if c["annee"] == y and c["niveau"] != "OK")
        if not y:
            row["statut"] = "Exercice non détecté"
        elif row.get("equilibre") == "ÉCART" or row["nb_anomalies"]:
            row["statut"] = "À revoir"
    except Exception as e:  # une liasse illisible ne bloque pas le lot
        row["statut"] = f"Erreur : {type(e).__name__}"
    return row


def summarize_dir(extr_dir):
    rows = []
    for f in sorted(Path(extr_dir).glob("*.json")):
        try:
            doc = json.loads(f.read_text(encoding="utf-8"))
        except Exception:
            rows.append({"fichier": f.name, "statut": "JSON illisible"}); continue
        if bv.is_v29d(doc):
            rows.append(summarize(f.stem, doc))
    return rows


def to_excel(rows, titre="Lot YaRisk"):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = "Synthèse"
    ws.append([c[1] for c in COLUMNS])
    head = PatternFill("solid", fgColor="00915A")
    for c in ws[1]:
        c.fill = head; c.font = Font(color="FFFFFF", bold=True); c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append([r.get(k) for k, _, _ in COLUMNS])
    for j, (k, lab, kind) in enumerate(COLUMNS, start=1):
        col = get_column_letter(j)
        ws.column_dimensions[col].width = 34 if k == "raison_sociale" else 28 if k == "fichier" else 18
        fmt_ = {"n": "#,##0", "p": "0.0%", "i": "0"}.get(kind)
        if fmt_:
            for cell in ws[col][1:]:
                cell.number_format = fmt_
    warn = PatternFill("solid", fgColor="FFF2CC")
    for i, r in enumerate(rows, start=2):
        if r.get("statut") != "OK":
            for c in ws[i]:
                c.fill = warn
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    info = wb.create_sheet("Informations")
    for line in [("Lot", titre), ("Généré le", datetime.now().strftime("%d/%m/%Y %H:%M")), ("Nombre de liasses", len(rows)),
                 ("Unité", "DZD"), ("Source", "YaRisk — moteur de lecture Mizan (liasses Série G) ou import JSON / CNRC"),
                 ("Statut « À revoir »", "écart actif / passif ou anomalie de formule ≥ 1 000 DZD : à vérifier avant usage"),
                 ("Usage", "Données de masse non revues par un analyste : statistiques, comparaisons et alertes uniquement ; "
                           "toute décision de crédit passe par la revue du dossier dans YaRisk.")]:
        info.append(line)
    info.column_dimensions["A"].width = 22; info.column_dimensions["B"].width = 110
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Lot de démonstration : liasses fictives cohérentes (quelques anomalies volontaires)
# ---------------------------------------------------------------------------
SECTEURS = ["Agroalimentaire", "BTP", "Négoce", "Pharmacie", "Transport", "Emballage", "Plasturgie", "Services informatiques",
            "Matériaux de construction", "Distribution automobile"]


def demo_docs(n=100, seed=7):
    rng = random.Random(seed)
    docs = {}
    for i in range(1, n + 1):
        sect = SECTEURS[i % len(SECTEURS)]
        name = f"Entreprise fictive {i:03d} — {sect}"
        nif = f"0999{rng.randint(10**10, 10**11 - 1)}"[:15]
        y = 2025 if rng.random() < 0.7 else 2024
        ca = rng.choice([0.08, 0.2, 0.5, 1.2, 3.0, 8.0]) * 1e9 * rng.uniform(0.6, 1.5)
        marge = rng.uniform(-0.04, 0.22)
        ebe = ca * marge
        dot = ca * rng.uniform(0.02, 0.05)
        ffin = ca * rng.uniform(0.005, 0.03)
        ibs = max(0.0, (ebe - dot - ffin) * 0.19)
        rn = ebe - dot - ffin - ibs
        conso = ca * rng.uniform(0.45, 0.75)
        pers = ca - conso - ebe
        stocks, clients = ca * rng.uniform(0.05, 0.25), ca * rng.uniform(0.08, 0.35)
        treso_a = ca * rng.uniform(0.01, 0.08)
        immo = ca * rng.uniform(0.2, 0.9)
        actif = immo + stocks + clients + treso_a
        capital = round(ca * rng.uniform(0.05, 0.2), -5)
        fp = actif * rng.uniform(0.15, 0.55)
        reserves = fp - capital - rn
        emprunts = actif * rng.uniform(0.05, 0.3)
        fourn = conso * rng.uniform(0.1, 0.3)
        autres = actif - fp - emprunts - fourn
        if autres < 0:
            emprunts += autres; autres = 0.0
        r = lambda v: float(round(v))
        A = {"autres_immobilisations_corporelles": {"net_n": r(immo)}, "total_actif_non_courant": {"net_n": r(immo)},
             "stocks_encours": {"net_n": r(stocks)}, "clients": {"net_n": r(clients)}, "tresorerie_actif": {"net_n": r(treso_a)},
             "total_actif_courant": {"net_n": r(stocks) + r(clients) + r(treso_a)}}
        A["total_general_actif"] = {"net_n": A["total_actif_non_courant"]["net_n"] + A["total_actif_courant"]["net_n"]}
        Pp = {"capital_emis": {"n": r(capital)}, "primes_reserves": {"n": r(reserves)}, "resultat_net_passif": {"n": r(rn)},
              "emprunts_dettes_financieres": {"n": r(emprunts)}, "fournisseurs_rattaches": {"n": r(fourn)}, "autres_dettes": {"n": r(autres)}}
        Pp["total_capitaux_propres"] = {"n": Pp["capital_emis"]["n"] + Pp["primes_reserves"]["n"] + Pp["resultat_net_passif"]["n"]}
        Pp["total_passifs_non_courants"] = {"n": Pp["emprunts_dettes_financieres"]["n"]}
        Pp["total_passifs_courants"] = {"n": Pp["fournisseurs_rattaches"]["n"] + Pp["autres_dettes"]["n"]}
        Pp["total_general_passif"] = {"n": Pp["total_capitaux_propres"]["n"] + Pp["total_passifs_non_courants"]["n"] + Pp["total_passifs_courants"]["n"]}
        side = lambda v: ({"n_credit": r(v)} if v >= 0 else {"n_debit": r(-v)})
        T = {"produits_fabriques": {"n_credit": r(ca)}, "chiffre_affaires_net": {"n_credit": r(ca)}, "production_exercice": {"n_credit": r(ca)},
             "matieres_premieres": {"n_debit": r(conso)}, "consommations_exercice": {"n_debit": r(conso)},
             "valeur_ajoutee_exploitation": side(r(ca) - r(conso)), "charges_personnel": {"n_debit": r(pers)},
             "excedent_brut_exploitation": side(r(ca) - r(conso) - r(pers)), "dotations_amortissements": {"n_debit": r(dot)},
             "resultat_operationnel": side(r(ca) - r(conso) - r(pers) - r(dot)), "charges_financieres": {"n_debit": r(ffin)},
             "resultat_financier": side(-r(ffin)), "resultat_ordinaire": side(r(ca) - r(conso) - r(pers) - r(dot) - r(ffin)),
             "impots_exigibles_resultats": {"n_debit": r(ibs)} if ibs else {},
             "resultat_net_exercice": side(r(ca) - r(conso) - r(pers) - r(dot) - r(ffin) - r(ibs))}
        T = {k: v for k, v in T.items() if v}
        Pp["resultat_net_passif"]["n"] = r(ca) - r(conso) - r(pers) - r(dot) - r(ffin) - r(ibs)
        Pp["primes_reserves"]["n"] = Pp["total_capitaux_propres"]["n"] - Pp["capital_emis"]["n"] - Pp["resultat_net_passif"]["n"]
        if i % 17 == 0:      # anomalie volontaire : une ligne du passif oubliée, total imprimé inchangé
            Pp.pop("fournisseurs_rattaches")
        if i % 23 == 0:      # anomalie volontaire : chiffres inversés dans un total
            s = str(int(A["total_actif_courant"]["net_n"])); A["total_actif_courant"]["net_n"] = float(s[:2] + s[3] + s[2] + s[4:])
        page = lambda tpl, table, rows, n_: {"page_id": f"p{n_:03d}", "numero_page_scannee": n_,
                                            "classification": {"page_template": tpl, "types": [table], "extraction_active": True},
                                            "entete_page": {"entreprise": name, "nif_15": nif}, "donnees": {table: rows}, "observations": []}
        docs[f"DEMO_{i:03d}_{y}"] = {"schema_version": "29.3", "document": {"fichier_source": f"DEMO_{i:03d}_{y}.pdf"},
                                     "identite": {"entreprise": name, "nif_15": nif, "annee": y, "nature": "FISCAL"},
                                     "pages": [page("ACTIF", "ACTIF", A, 1), page("PASSIF", "PASSIF", Pp, 2), page("TCR_FULL", "TCR", T, 3)],
                                     "controles": {}}
    return docs
