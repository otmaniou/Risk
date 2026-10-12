"""Ya-Risk — génère bilan_grid/model_saisie.js à partir du modèle Excel ANALYSTE (V20 / « ameliorer »).

À lancer une seule fois (ou à chaque nouvelle version du modèle) :
    python tools/build_saisie_model.py "MODEL ETAT FINANCIER ANALYSTE.xlsx"

Le fichier produit contient, pour les feuilles Saisie actif / Saisie passif / Saisie TCR / Saisie / Etatsfin :
  - les formules EXACTES du modèle (évaluées dans le navigateur par grid.js) ;
  - les styles d'affichage de Saisie et Etatsfin (couleurs, gras, formats de nombre) ;
  - la « préparation » de la feuille Saisie, identique dans le web et dans l'Excel généré :
      * les valeurs d'exemple du modèle (autre société) sont vidées et deviennent des cellules de saisie ;
      * chaque ligne reprend, pour toutes les années, la formule de la colonne I (exercice N) décalée vers
        l'exercice correspondant (Saisie actif : 3 colonnes par exercice, Saisie passif : 1, Saisie TCR : 2) ;
      * les cellules cassées du modèle (#REF!, liens vers d'autres classeurs) sont réparées quand la ligne
        fournit le motif ; les autres restent en erreur et sont listées.
"""
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string as cidx, get_column_letter as cl

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "bilan_grid" / "model_saisie.js"
CRF_OUT = ROOT / "bilan_grid" / "crf_template.xlsx"

CFG = {
    "ACTIF_GROUPES": {"N": ["P", "Q", "R"], "N-1": ["M", "N", "O"], "N-2": ["J", "K", "L"], "N-3": ["G", "H", "I"], "N-4": ["D", "E", "F"]},
    "PASSIF_COLS": {"N": "H", "N-1": "G", "N-2": "F", "N-3": "E", "N-4": "D"},
    "TCR_PAIRES": {"N": ["M", "N"], "N-1": ["K", "L"], "N-2": ["I", "J"], "N-3": ["G", "H"], "N-4": ["E", "F"]},
    "ACTIF_ROWS": {"ecarts_acquisition_goodwill": 10, "immobilisations_incorporelles": 11, "immobilisations_corporelles": 12, "terrains": 13,
                   "batiments": 14, "autres_immobilisations_corporelles": 15, "immobilisations_en_concession": 16, "immobilisations_en_cours": 17,
                   "immobilisations_financieres": 18, "titres_mis_en_equivalence": 19, "autres_participations_creances": 20, "autres_titres_immobilises": 21,
                   "prets_actifs_financiers_non_courants": 22, "impots_differes_actif": 23, "total_actif_non_courant": 24, "stocks_encours": 26,
                   "creances_emplois_assimiles": 27, "clients": 28, "autres_debiteurs": 29, "impots_assimiles_actif": 30, "autres_creances_assimiles": 31,
                   "disponibilites_assimiles": 32, "placements_financiers_courants": 33, "tresorerie_actif": 34, "total_actif_courant": 35, "total_general_actif": 36},
    "PASSIF_ROWS": {"capital_emis": 6, "capital_non_appele": 7, "primes_reserves": 8, "ecart_reevaluation": 9, "ecart_equivalence": 10, "resultat_net_passif": 11,
                    "report_a_nouveau": 12, "part_societe_consolidante": 13, "part_minoritaires": 14, "total_capitaux_propres": 15,
                    "emprunts_dettes_financieres": 17, "impots_differes_provisionnes": 18, "autres_dettes_non_courantes": 19, "provisions_produits_avance": 20,
                    "total_passifs_non_courants": 21, "fournisseurs_rattaches": 23, "impots_passif": 24, "autres_dettes": 25, "tresorerie_passif": 26,
                    "total_passifs_courants": 27, "total_general_passif": 28},
    "TCR_ROWS": {"ventes_marchandises": 6, "produits_fabriques": 7, "prestations_services": 8, "ventes_travaux": 9, "produits_annexes": 10,
                 "rabais_remises_ristournes_accordes": 11, "chiffre_affaires_net": 12, "production_stockee_destockee": 13, "production_immobilisee": 14,
                 "subvention_exploitation": 15, "production_exercice": 16, "achats_marchandises_vendues": 17, "matieres_premieres": 18,
                 "autres_approvisionnements": 19, "variation_stocks": 20, "achats_etudes_prestations": 21, "autres_consommations": 22,
                 "rabais_remises_ristournes_obtenus_achats": 23, "sous_traitance_generale": 24, "locations": 25, "entretien_reparations": 26,
                 "primes_assurances": 27, "personnel_exterieur": 28, "remuneration_intermediaires": 29, "publicite": 30, "deplacements_missions": 31,
                 "autres_services": 32, "rabais_remises_ristournes_obtenus_services": 33, "consommations_exercice": 34, "valeur_ajoutee_exploitation": 35,
                 "charges_personnel": 36, "impots_taxes_assimiles": 37, "excedent_brut_exploitation": 38, "autres_produits_operationnels": 39,
                 "autres_charges_operationnelles": 40, "dotations_amortissements": 41, "provisions": 42, "pertes_valeur": 43,
                 "reprises_pertes_valeur_provisions": 44, "resultat_operationnel": 45, "produits_financiers": 46, "charges_financieres": 47,
                 "resultat_financier": 48, "resultat_ordinaire": 49, "elements_extraordinaires_produits": 50, "elements_extraordinaires_charges": 51,
                 "resultat_extraordinaire": 52, "impots_exigibles_resultats": 53, "impots_differes_resultats": 54, "resultat_net_exercice": 55},
}
RANKS = ["N", "N-1", "N-2", "N-3", "N-4"]
SHEETS = ["Saisie actif", "Saisie passif", "Saisie TCR", "Saisie", "Etatsfin"]

# Colonnes d'exercice : rang 0 = N (dernier exercice), 1 = N-1 ...  (Saisie D = N-5, sert de base aux variations)
SAISIE_YEAR_COLS = ["I", "H", "G", "F", "E", "D"]
ETATS_YEAR_COLS = [["H", "N"], ["G", "M"], ["F", "L"], ["E", "K"], ["D", "J"], ["C", None]]
# Décalage de colonnes par exercice dans les feuilles de saisie liasse, et plage valide
LINK_SHIFT = {"Saisie actif": (3, "D", "R"), "Saisie passif": (1, "D", "H"), "Saisie TCR": (2, "E", "N")}

SAISIE_LAST_ROW = 234
SAISIE_INPUT_ROWS = [r for r in range(4, 193) if r not in (114, 177)] + [215]
SAISIE_HELPER_SRC = {"R119": None, "R120": None, "R122": None}   # aides du contrôle Actif = Passif (ligne 118)
ETATS_INPUTS = [f"{c}{r}" for r in (97, 98, 99) for c in "JKLMN"]   # échéancier de la dette MLT (2 ans, 2/5 ans, > 5 ans)

STR_RE = re.compile(r'"(?:[^"]|"")*"')
REF_RE = re.compile(r"(?P<pre>(?:'(?P<qs>[^']+)'|(?P<us>[A-Za-z_\[\]0-9][\w\.\[\]]*))!)?(?P<cd>\$?)(?P<c>[A-Z]{1,3})(?P<rd>\$?)(?P<r>\d+)(?![\w(])")


def has_refs(f):
    return isinstance(f, str) and f.startswith("=") and bool(REF_RE.search(STR_RE.sub('""', f))) and "#REF" not in f


def translate(formula, delta, home):
    """Décale une formule de `delta` colonnes d'exercice. Renvoie None si une référence sort des colonnes valides."""
    parts, pos, ok = [], 0, [True]
    prev = {"sheet": None, "end": -1}

    def fix(seg, offset):
        def rep(m):
            sheet = m.group("qs") or m.group("us")
            start = offset + m.start()
            if sheet is None and prev["end"] == start - 1 and prev["sheet"]:
                sheet = prev["sheet"]   # second terme d'une plage 'Feuille'!A1:B2
            prev["sheet"] = sheet; prev["end"] = offset + m.end()
            sh = (sheet or home).strip("'")
            if m.group("cd") == "$" or delta == 0:
                return m.group(0)
            mult, lo, hi = LINK_SHIFT.get(sh, (1, "A", "XFD"))
            nc = cidx(m.group("c")) + delta * mult
            if nc < cidx(lo) or nc > cidx(hi):
                ok[0] = False
                return m.group(0)
            return (m.group("pre") or "") + m.group("cd") + cl(nc) + m.group("rd") + m.group("r")
        return REF_RE.sub(rep, seg)

    for m in STR_RE.finditer(formula):
        parts.append(fix(formula[pos:m.start()], pos)); parts.append(m.group(0)); pos = m.end()
    parts.append(fix(formula[pos:], pos))
    return "".join(parts) if ok[0] else None


def jval(v):
    if isinstance(v, datetime):
        return {"d": v.strftime("%Y-%m-%d")}
    if isinstance(v, date):
        return {"d": v.isoformat()}
    if isinstance(v, (int, float)) or isinstance(v, str) or v is None or isinstance(v, bool):
        return v
    return str(v)


def color_of(c):
    try:
        if c is not None and c.type == "rgb" and isinstance(c.rgb, str) and len(c.rgb) == 8:
            return "#" + c.rgb[2:]
    except Exception:
        pass
    return None



# ---------------------------------------------------------------------------------------------------------------
# CRF (BNP Paribas Credit Rating Framework) : feuilles de restitution reliées à la liasse (tools/crf_mapping.py)
# ---------------------------------------------------------------------------------------------------------------
TOK_RE = re.compile(r"\b(A|P|T):([a-z_0-9]+)(?::(brut|amort|net))?(@1)?")


def crf_formula(expr, k):
    """Expression de correspondance -> formule Excel (Million DZD) pour l'exercice de rang k (0 = N)."""
    bad = [False]

    def rep(m):
        kind, rid, part, prev = m.groups()
        kk = k + (1 if prev else 0)
        if kk >= len(RANKS):
            bad[0] = True; return "0"
        rk = RANKS[kk]
        if kind == "A":
            return f"'Saisie actif'!{CFG['ACTIF_GROUPES'][rk][['brut', 'amort', 'net'].index(part)]}{CFG['ACTIF_ROWS'][rid]}"
        if kind == "P":
            return f"'Saisie passif'!{CFG['PASSIF_COLS'][rk]}{CFG['PASSIF_ROWS'][rid]}"
        d, c = CFG["TCR_PAIRES"][rk]; row = CFG["TCR_ROWS"][rid]
        return f"('Saisie TCR'!{c}{row}-'Saisie TCR'!{d}{row})"

    e = TOK_RE.sub(rep, expr)
    return None if bad[0] else f'=IF(ROUND({e},3)=0,"",ROUND({e},3)/1000)'


def crf_human(expr):
    sys.path.insert(0, str(ROOT))
    import bilan_view as bv
    lab = {"A": {r[1]: r[0] for r in bv.ACTIF_LAYOUT}, "P": {r[1]: r[0] for r in bv.PASSIF_LAYOUT},
           "T": {r[2]: (r[1] or r[0]) for r in bv.TCR_LAYOUT}}
    part = {"brut": "Brut", "amort": "Amort.", "net": "Net"}

    def rep(m):
        kind, rid, p, prev = m.groups()
        t = {"A": "Actif", "P": "Passif", "T": "TCR"}[kind] + " › " + str(lab[kind].get(rid, rid)).strip()
        if p: t += " (" + part[p] + ")"
        return "[" + t + (" N-1" if prev else "") + "]"
    return TOK_RE.sub(rep, expr)


def build_crf(crf_xlsx, cells, display, style_of):
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import crf_mapping as cm
    for (sh, r), (expr, lvl, note) in cm.MAPPING.items():       # contrôle des noms de rubriques
        for m in TOK_RE.finditer(expr):
            rows = {"A": CFG["ACTIF_ROWS"], "P": CFG["PASSIF_ROWS"], "T": CFG["TCR_ROWS"]}[m.group(1)]
            assert m.group(2) in rows, (sh, r, m.group(2))
    wb = openpyxl.load_workbook(crf_xlsx)
    inputs, mapped, text, headers, created, doc = [], [], [], [], [], []
    for name in cm.SHEETS:
        ws = wb[name]
        cols = cm.YEAR_COLS.get(name)
        # en-tête : date de création, analyste (saisi une fois sur Bilan - Actifs), dates d'exercice
        for c in ws[1]:
            if c.value == "Créé le":
                created.append([name, ws.cell(1, c.column + 1).coordinate]); ws.cell(1, c.column + 1).value = None
            if c.value == "Analyste :":
                tgt = ws.cell(1, c.column + 1)
                if name == cm.AB:
                    tgt.value = None; text.append(f"{name}!{tgt.coordinate}")
                else:
                    tgt.value = "=IF('Bilan - Actifs'!E1=\"\",\"\",'Bilan - Actifs'!E1)"
        if cols:
            for k, col in enumerate(cols):
                for r in (5, 6, 7):
                    ws[f"{col}{r}"].value = None; headers.append([name, f"{col}{r}", k, r])
            skip = cm.NO_INPUT_ROWS.get(name, set())
            for r in range(15, ws.max_row + 1):
                if ws[f"A{r}"].value in (None, "") or r in skip:
                    continue
                m = cm.MAPPING.get((name, r))
                if m:
                    doc.append({"feuille": name, "ligne": r, "rubrique": str(ws[f"A{r}"].value).strip(), "liasse": crf_human(m[0]),
                                "niveau": m[1], "note": m[2]})
                for k, col in enumerate(cols):
                    a = f"{col}{r}"; v = ws[a].value
                    if name in (cm.AB, cm.PB, cm.CR, cm.TF) and ws[a].number_format == "General":
                        ws[a].number_format = "#,##0.00"      # affichage homogène des montants (Million DZD)
                    if (name, a) in cm.CELL_FORMULAS:
                        ws[a].value = cm.CELL_FORMULAS[(name, a)]; continue
                    if isinstance(v, str) and v.startswith("="):
                        continue
                    ws[a].value = None; inputs.append(f"{name}!{a}")
                    if ws[a].number_format == "General":
                        ws[a].number_format = "#,##0.00"
                    if m:
                        f = crf_formula(m[0], k)
                        if f:
                            ws[a].value = f; mapped.append(f"{name}!{a}")
        else:   # Leveraged Analysis & ECB : réponses qualitatives
            for r in range(3, ws.max_row + 1):
                if ws[f"E{r}"].value is not None:
                    ws[f"E{r}"].value = None; text.append(f"{name}!E{r}")
        d = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    d[c.coordinate] = jval(c.value)
        cells[name] = d
        st = {}
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=10):
            for c in row:
                st[c.coordinate] = style_of(c)
        display[name] = {"styles": st, "widths": {k: round(v.width, 2) for k, v in ws.column_dimensions.items() if v.width},
                         "hiddenRows": [], "lastRow": ws.max_row}
    wb.save(CRF_OUT)
    return {"sheets": cm.SHEETS, "yearCols": {k: list(v) for k, v in cm.YEAR_COLS.items()}, "inputs": inputs, "mapped": mapped,
            "text": text, "headers": headers, "created": created, "doc": doc, "template": CRF_OUT.name, "source": Path(crf_xlsx).name}


def main(xlsx, crf_xlsx=None):
    wb = openpyxl.load_workbook(xlsx)
    cells = {}
    for name in SHEETS:
        ws = wb[name]; d = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    d[c.coordinate] = jval(c.value)
        cells[name] = d
    # --- 0. Règles Ya-Risk imposées au modèle (identiques à l'écran, à l'Excel et aux contrôles du notebook) ---
    #   immobilisations corporelles = terrains + bâtiments + autres + concession + en cours (lignes 13 à 17)
    #   actif non courant = goodwill + incorporelles + corporelles + financières (lignes 10, 11, 12, 18)
    model_fixes = []
    for c in "DEFGHIJKLMNOPQR":
        model_fixes += [["Saisie actif", f"{c}12", f"=SUM({c}13:{c}17)"], ["Saisie actif", f"{c}9", f"=+{c}10+{c}11+{c}12+{c}18"]]
    # Saisie reprenait « corporelles + en cours » : en cours est désormais dans les corporelles
    model_fixes += [["Saisie", "I14", "=+'Saisie actif'!P12"], ["Saisie", "I16", "=+'Saisie actif'!Q12"]]
    for sh, a, f in model_fixes:
        cells[sh][a] = f
    original = json.loads(json.dumps(cells))
    for sh, a, f in model_fixes:      # le patch Saisie doit aussi porter ces 2 corrections
        if sh == "Saisie":
            original[sh][a] = wb[sh][a].value

    # --- 1. Zones de saisie liasse vidées (comme clear_writable_zones du notebook) ---
    def is_f(v): return isinstance(v, str) and v.startswith("=")
    liasse_inputs = {"Saisie actif": [], "Saisie passif": [], "Saisie TCR": []}
    for rank in RANKS:
        for rw in CFG["ACTIF_ROWS"].values():
            for c in CFG["ACTIF_GROUPES"][rank]:
                a = f"{c}{rw}"
                if not is_f(cells["Saisie actif"].get(a)): cells["Saisie actif"].pop(a, None); liasse_inputs["Saisie actif"].append(a)
        for rw in CFG["PASSIF_ROWS"].values():
            a = f"{CFG['PASSIF_COLS'][rank]}{rw}"
            if not is_f(cells["Saisie passif"].get(a)): cells["Saisie passif"].pop(a, None); liasse_inputs["Saisie passif"].append(a)
        for rw in CFG["TCR_ROWS"].values():
            for c in CFG["TCR_PAIRES"][rank]:
                a = f"{c}{rw}"
                if not is_f(cells["Saisie TCR"].get(a)): cells["Saisie TCR"].pop(a, None); liasse_inputs["Saisie TCR"].append(a)

    # --- 2. Préparation de Saisie ---
    S = cells["Saisie"]; repairs = []; inputs = []
    order = ["I", "H", "G", "F", "E", "D"]
    for r in range(3, 219):
        if r == 3:
            continue
        canon = None
        for col in order:
            v = S.get(f"{col}{r}")
            if has_refs(v):
                canon = (col, v); break
        if canon:
            src_col, f = canon
            for col in "DEFGHI":
                a = f"{col}{r}"
                nf = translate(f, cidx(col) - cidx(src_col), "Saisie")
                if nf is None:
                    if a in S: S.pop(a)
                    if r in SAISIE_INPUT_ROWS: inputs.append(a)
                    continue
                old = S.get(a)
                if old != nf:
                    if isinstance(old, str) and ("#REF" in old or "[" in old):
                        repairs.append({"feuille": "Saisie", "cellule": a, "avant": old, "apres": nf})
                    S[a] = nf
        elif r in SAISIE_INPUT_ROWS and (S.get(f"B{r}") or S.get(f"C{r}")):
            for col in "DEFGHI":
                S.pop(f"{col}{r}", None); inputs.append(f"{col}{r}")
    # aides de la ligne 118 pour les colonnes E et D (colonnes N et M)
    for src in ("R119", "R120", "R122"):
        for tgt in ("N", "M"):
            a = tgt + src[1:]
            S[a] = translate(original["Saisie"][src], cidx(tgt) - cidx("R"), "Saisie")
    # En-têtes : dates des exercices (écrites dynamiquement), société, groupe
    for col in "DEFGHI":
        S.pop(f"{col}3", None)
    S.pop("C2", None); S.pop("C3", None)

    # --- 3. Etatsfin : réparation des #REF! d'après la ligne ---
    E = cells["Etatsfin"]
    for a, v in list(E.items()):
        if isinstance(v, str) and "#REF" in v:
            m = re.match(r"([A-Z]+)(\d+)", a); col, r = m.group(1), int(m.group(2))
            group = "CDEFGH" if col in "CDEFGH" else "JKLMN"
            for sc in reversed(group[: group.index(col)]):
                sv = E.get(f"{sc}{r}")
                if has_refs(sv):
                    nf = translate(sv, cidx(col) - cidx(sc), "Etatsfin")
                    if nf:
                        repairs.append({"feuille": "Etatsfin", "cellule": a, "avant": v, "apres": nf}); E[a] = nf
                    break
    for a in ETATS_INPUTS:
        E.pop(a, None)

    broken = []
    for name in SHEETS:
        for a, v in cells[name].items():
            if isinstance(v, str) and v.startswith("=") and ("#REF" in v or re.search(r"\[\d+\]", v)):
                broken.append({"feuille": name, "cellule": a, "formule": v})

    # --- 4. Patch Excel (Saisie / Etatsfin) : différences modèle -> préparé ---
    patch = []
    for name in ("Saisie", "Etatsfin"):
        keys = set(original[name]) | set(cells[name])
        for a in sorted(keys, key=lambda x: (int(re.sub(r"\D", "", x)), cidx(re.sub(r"\d", "", x)))):
            if original[name].get(a) != cells[name].get(a):
                patch.append([name, a, cells[name].get(a)])

    # --- 5. Styles d'affichage ---
    styles, sidx = [], {}

    def style_of(c):
        f = c.font; fill = c.fill
        bg = color_of(fill.fgColor) if fill is not None and fill.fill_type == "solid" else None
        st = {"bg": bg, "fg": color_of(f.color) if f is not None else None, "b": bool(f.b), "i": bool(f.i),
              "sz": float(f.sz or 10), "h": c.alignment.horizontal, "w": bool(c.alignment.wrap_text), "nf": c.number_format or "General",
              "bt": c.border.top.style, "bb": c.border.bottom.style}
        key = json.dumps(st, sort_keys=True)
        if key not in sidx:
            sidx[key] = len(styles); styles.append(st)
        return sidx[key]

    display = {}
    for name, last_row, last_col in (("Saisie", SAISIE_LAST_ROW, "J"), ("Etatsfin", 102, "N")):
        ws = wb[name]; st = {}
        for row in ws.iter_rows(min_row=1, max_row=last_row, max_col=cidx(last_col)):
            for c in row:
                st[c.coordinate] = style_of(c)
        display[name] = {
            "styles": st,
            "widths": {k: round(v.width, 2) for k, v in ws.column_dimensions.items() if v.width},
            "hiddenRows": [r for r, d in ws.row_dimensions.items() if d.hidden or (d.height is not None and d.height < 3)],
            "lastRow": last_row,
        }

    crf = build_crf(crf_xlsx, cells, display, style_of) if crf_xlsx else None

    model = {
        "crf": crf,
        "source": Path(xlsx).name, "built": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "cfg": CFG, "ranks": RANKS, "cells": cells, "display": display, "styles": styles,
        "liasseInputs": liasse_inputs,
        "saisieInputs": sorted(set(inputs), key=lambda x: (int(x[1:]), x[0])),
        "etatsInputs": ETATS_INPUTS,
        "saisieYearCols": SAISIE_YEAR_COLS, "etatsYearCols": ETATS_YEAR_COLS,
        "patch": patch, "repairs": repairs, "broken": broken,
        "modelFixes": [x for x in model_fixes if x[0] == "Saisie actif"],
    }
    if crf:
        print(f"CRF : {len(crf['inputs'])} cellules de saisie dont {len(crf['mapped'])} reliées à la liasse | modèle nettoyé {CRF_OUT.name}")
    OUT.write_text("window.YR_MODEL=" + json.dumps(model, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
    print(f"OK {OUT} — {OUT.stat().st_size // 1024} Ko | patch {len(patch)} | réparées {len(repairs)} | encore cassées {len(broken)} | saisies {len(inputs)}")
    for b in broken:
        print("  cassée :", b["feuille"], b["cellule"], b["formule"][:90])
    for x in repairs:
        print("  réparée:", x["feuille"], x["cellule"])


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
